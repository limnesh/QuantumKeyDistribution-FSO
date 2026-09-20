"""Write audit/provenance and append documentation without replacing baselines."""
from pathlib import Path
import ast,hashlib,json,re,zipfile
from docx import Document
ROOT=Path(__file__).resolve().parents[1]

def main():
    original=json.loads((ROOT/'validation/original_inventory.json').read_text())
    source_inventory=[]
    for row in original:
        p=ROOT/row['path']
        if p.suffix not in ('.m','.py','.ipynb'): continue
        text=p.read_text(encoding='utf-8')
        item=dict(path=row['path'])
        if p.suffix=='.py':
            tree=ast.parse(text); item['definitions']=[n.name for n in ast.walk(tree) if isinstance(n,(ast.FunctionDef,ast.ClassDef))]
            item['imports']=[ast.unparse(n) for n in ast.walk(tree) if isinstance(n,(ast.Import,ast.ImportFrom))]
            item['calls']=sorted(set(ast.unparse(n.func) for n in ast.walk(tree) if isinstance(n,ast.Call)))
        elif p.suffix=='.m':
            item['definitions']=re.findall(r'^function\s+(.+)',text,re.M)
            item['calls']=sorted(set(re.findall(r'\b([A-Za-z]\w*)\s*\(',text)))
        else:
            nb=json.loads(text); item['cells']=len(nb['cells'])
            item['section_headings']=[str(c['source'])[:200] for c in nb['cells'] if c['cell_type']=='markdown']
        source_inventory.append(item)
    (ROOT/'validation/dependency_map.json').write_text(json.dumps(source_inventory,indent=2),encoding='utf-8')
    audit='''# Repository audit and dependency map

Audit preceded implementation. `git status --short` returned `fatal: not a git repository (or any of the parent directories): .git`.
No reset or deletion was performed. Original sources, notebooks, HTML and unlocked report documents are retained in
`validation/pre_extension_backup.zip`; original SHA-256 identities are in `validation/original_inventory.json`.
Historical `example_output`, `results`, and comparison folders were not reused for extension exports.

## Mathematical and execution dependencies

| Area | Existing source and callers | Extension |
|---|---|---|
| Stage 1 | Notebook inline optical/BB84/legacy DPS functions; qkd_link_model -> qkd_simulate -> dashboard/draw/export; qkd_bb84_eve, qkd_ldpc_graph/decode, qkd_toeplitz_hash | shared.physics / shared.scenarios.stage1; extension_simulate(1); companion dashboard |
| Stage 2 | leo_model link_at_elevation/qkd_probabilities -> simulate_pass; leo_link_model -> leo_simulate; leo_dashboard/draw/export; notebook builders | same geometry, shared detector/fading adapter stage2 / extension_simulate(2) |
| Stage 3A | network_model -> stage-2 leo_model; simulate_network -> widest_path; stage3_model trusted_inventory; Octave network_simulate/network_widest_path/stage3_compare | per-edge protocol/fading -> routing -> existing classical inventory semantics |
| Stage 3B | stage3_model.bbm92_rates and bbm92_rates.m; stage3_compare and stage3 reports | both faded arms passed into original singles/coincidence/noise equations |
| Stage 4 | local copies of network_model, leo_model, stage3_model; network_setup_path finds stage 2 in Octave | shared.scenarios.network; independent ground-edge streams; vacuum unchanged |
| Entry points | START_PROJECT.m -> four START_HERE.m files -> qkd/leo/network dashboards | new menu and RUN_DPS_TURBULENCE.m in every folder |
| Reports | qkd/leo/network_export_results.m, network_report.py, stage3_report.py, notebook builders | shared.reporting + extension_export/study; CSV/JSON/HTML/PNG to distinct extension folders |

There was no `shared/` folder at audit. Stage 3/4 carried copied Python and Octave modules.
New equations have one Python and one Octave implementation. Compatibility copies and historical
notebook mathematics remain intact; no broad reorganization was attempted.

## Legacy calculations

Stage 1 `qkd_link_model.m` and notebook Sections 5–6 contain the visibility-only DPS teaching proxy:
gain times a BB84-like entropy fraction, without DPS pulse-slot operation. It remains historical.
Stage 1 `qkd_simulate.m` and notebook Section 7 contain the independent 5,000-sample log-normal
teaching knob sigma=strength*sqrt(distance_km/10). The new sigma is directly log-irradiance
standard deviation; it has no claimed Cn² mapping. Stage 2/3/4 baselines had no stochastic fading.
BB84 variants use the existing additive small-background convention in Stage 1 and exclusive
background-only gates in satellite stages. The original BBM92 approximation guards remain active.

## References and document handling

All repository PDF paths were searched. Neither of the two requested PDF names/titles was present.
The online tutorial Section 5.4 was consulted; the weak-turbulence full text was unavailable.
See DPS_THEORY.md for exact consulted sources and the model decision. A Word lock exists for
`QKD - FSO Report.docx`; that working report is untouched. The unlocked
`FSO_QKD_Detailed_Project_Report.docx` receives an appendix using existing styles.
`DPS_REPORT_INSERTION.md` is available for the working report.

## Inspected source inventory

The machine-readable definition/import/call inventory is `validation/dependency_map.json`.
The paths below identify actual source, test, launcher, export and notebook files read during the audit and implementation.

'''
    audit+='\n'.join('- `'+x['path'].replace('\\','/')+'`' for x in source_inventory)+'\n'
    (ROOT/'DPS_REPOSITORY_AUDIT.md').write_text(audit,encoding='utf-8')
    note='''
## DPS and atmospheric turbulence extension (September 2026)

The earlier sections describe the preserved historical baselines. The new shared extension covers
Stages 1, 2, trusted 3A, entanglement-based 3B and network 4. DPS outputs are **illustrative rate
proxies**, never certified secret keys. BBM92 remains a separate protocol.

- Theory, equations and security limitations: [DPS_THEORY.md]({root}DPS_THEORY.md).
- Implementation and executed verification: [implementation report]({root}DPS_TURBULENCE_IMPLEMENTATION_REPORT.md).
- Python: from the project root, `python tools/run_dps_studies.py`.
- Octave: use **DPS / atmospheric fading** in the existing dashboard menu, or run this scenario's `RUN_DPS_TURBULENCE.m`.
- All three fading modes are supported; library default is `none`. Demonstration reports explicitly use correlated fading.
- Outputs use separate `dps_turbulence_results` and `dps_turbulence_notebook_results` folders.
- Re-execute notebooks and HTML: `python tools/rebuild_notebooks.py` from the project root.

'''
    for p in [ROOT/'README.md',ROOT/'SCIENTIFIC_AUDIT.md',ROOT/'VERIFICATION.md',*ROOT.glob('[1234]-*/README.md'),ROOT/'1-Basic LOS FSO/QKD_FSO_Octave/README.md',ROOT/'1-Basic LOS FSO/study_guide_source/expanded_guide.md']:
        marker='## DPS and atmospheric turbulence extension (September 2026)'
        text=p.read_text(encoding='utf-8')
        if marker in text: text=text.split(marker)[0].rstrip()+'\n'
        depth=len(p.relative_to(ROOT).parts)-1
        p.write_text(text+note.format(root='../'*depth),encoding='utf-8')
    index=ROOT/'PROJECT_INDEX.html'; text=index.read_text(encoding='utf-8'); marker='<!-- DPS_EXTENSION_INDEX -->'
    text=re.sub(re.escape(marker)+r'.*?</section>','',text,flags=re.S)
    block='\n'+marker+'\n<section style="margin:2rem"><h2>DPS and atmospheric turbulence</h2><p>DPS rates and derived relay budgets are illustrative proxies. BBM92 remains separate.</p><p><a href="DPS_TURBULENCE_IMPLEMENTATION_REPORT.md">Implementation and verification</a> · <a href="DPS_THEORY.md">Scientific specification</a></p><ul>'
    from urllib.parse import quote
    for stage in range(1,5):
        folder=next(ROOT.glob(f'{stage}-*'))
        block+=f'<li><a href="{quote(folder.name)}/dps_turbulence_results/report.html">Stage {stage}: executed DPS / fading analysis</a></li>'
    block+='</ul></section>\n'
    text=text.replace('</body>',block+'</body>') if '</body>' in text else text+block
    index.write_text(text,encoding='utf-8')
    theory=(ROOT/'DPS_THEORY.md').read_text(encoding='utf-8')
    (ROOT/'DPS_REPORT_INSERTION.md').write_text('# Appendix D — DPS and atmospheric turbulence\n\n'+theory,encoding='utf-8')
    docpath=ROOT/'FSO_QKD_Detailed_Project_Report.docx'; doc=Document(docpath)
    title='Appendix D. DPS QKD and atmospheric turbulence extension'
    if not any(p.text==title for p in doc.paragraphs):
        original_paragraphs=[(p.text,p.style.name) for p in doc.paragraphs]; original_tables=len(doc.tables)
        doc.add_page_break(); doc.add_heading(title,level=1)
        doc.add_paragraph('This appendix supplements the historical results without renumbering existing figures. Generated data are simulated, not measured. See DPS_TURBULENCE_IMPLEMENTATION_REPORT.md for executed test evidence.')
        for block in theory.split('\n\n'):
            if block.startswith('# '): continue
            if block.startswith('## '): doc.add_heading(block[3:].replace('\n',' '),level=2)
            else: doc.add_paragraph(block.replace('```','').replace('**','').replace('`','').replace('\n',' '))
        doc.save(docpath)
        reread=Document(docpath)
        assert [(p.text,p.style.name) for p in reread.paragraphs[:len(original_paragraphs)]]==original_paragraphs
        assert len(reread.tables)==original_tables
        (ROOT/'validation/docx_preservation.json').write_text(json.dumps(dict(status='passed',original_paragraphs=len(original_paragraphs),original_tables=original_tables,appended_paragraphs=len(reread.paragraphs)-len(original_paragraphs))),encoding='utf-8')

if __name__=='__main__': main()
